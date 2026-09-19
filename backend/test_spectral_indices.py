import json
import os
import shutil
import sys
import tempfile
import unittest
import uuid
from pathlib import Path
from unittest.mock import patch

import numpy as np
import rasterio
from openpyxl import load_workbook
from rasterio.transform import from_origin

sys.path.append(os.path.join(os.path.dirname(__file__), "."))

from applications import create_app
from applications.common.path_global import generate_url, up_url
from applications.extensions import db
from applications.kml_roi.index_sync import sync_miner_index_rows
from applications.kml_roi.kml import load_kml_features
from applications.kml_roi.kml_merge import merge_kml_increment
from applications.models.analysis import Analysis
from applications.models.project import ProjectMineBinding
from applications.models.project_spatial import ProjectSpatialResource
from applications.project_hub.service import create_project


class TestSpectralIndicesAPI(unittest.TestCase):
    def setUp(self):
        self.app = create_app("testing")
        self.app.config["PROPAGATE_EXCEPTIONS"] = True
        self.client = self.app.test_client()
        self.ctx = self.app.app_context()
        self.ctx.push()
        db.create_all()
        self.login_as_admin()

        # 测试隔离（AGENTS §8）：上传/生成目录指向临时目录，不写真实仓库
        # static/upload/；范式与 test_large_tiff_upload 一致
        from applications.extensions.flask_uploads import UploadConfiguration
        from applications.extensions.init_upload import IMAGES_WITH_TIFF

        self.upload_tmp = tempfile.mkdtemp(prefix="spectral-isolated-")
        self.up_dir = self.upload_tmp
        self.generate_dir = os.path.join(self.upload_tmp, "res")
        os.makedirs(self.generate_dir, exist_ok=True)
        self.app.config["UPLOADED_PHOTOS_DEST"] = self.upload_tmp
        self.app.upload_set_config["photos"] = UploadConfiguration(
            self.upload_tmp, None, IMAGES_WITH_TIFF, ()
        )
        from applications.api import analysis as analysis_module

        for target, value in (
            ("up_dir", self.up_dir),
            ("generate_dir", self.generate_dir),
        ):
            patcher = patch.object(analysis_module, target, value)
            patcher.start()
            self.addCleanup(patcher.stop)
        self.addCleanup(shutil.rmtree, self.upload_tmp, ignore_errors=True)

        self.created_upload_files = []
        self.created_result_files = []
        self.temp_dirs = []

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.ctx.pop()

        for fpath in self.created_upload_files:
            if os.path.exists(fpath):
                os.remove(fpath)

        for fpath in self.created_result_files:
            if os.path.exists(fpath):
                os.remove(fpath)

        for dpath in self.temp_dirs:
            if os.path.exists(dpath):
                shutil.rmtree(dpath)

    def sync_admin(self, password="Secret123!"):
        os.environ["ADMIN_USERNAME"] = "admin"
        os.environ["ADMIN_PASSWORD"] = password
        self.app.config["ADMIN_USERNAME"] = "admin"
        self.app.config["ADMIN_PASSWORD"] = password
        from applications.auth.service import sync_admin_from_env

        return sync_admin_from_env()

    def login_as_admin(self, password="Secret123!"):
        self.sync_admin(password)
        response = self.client.post(
            "/api/auth/login",
            json={"username": "admin", "password": password},
        )
        self.assertEqual(response.status_code, 200)
        return response

    def _make_tif(self, band_count, basename_prefix="spec"):
        fname = f"{basename_prefix}_{uuid.uuid4().hex[:8]}.tif"
        fpath = os.path.join(self.up_dir, fname)
        height, width = 16, 16
        transform = from_origin(0, 0, 1, 1)

        base = np.linspace(10, 200, num=height * width, dtype=np.float32).reshape(height, width)
        bands = []
        for i in range(band_count):
            bands.append(base + (i * 3))

        arr = np.stack(bands, axis=0)
        with rasterio.open(
            fpath,
            "w",
            driver="GTiff",
            height=height,
            width=width,
            count=band_count,
            dtype=np.float32,
            transform=transform,
        ) as dst:
            for i in range(band_count):
                dst.write(arr[i], i + 1)

        self.created_upload_files.append(fpath)
        return fname

    def _make_constant_geo_tif(self, basename_prefix="geo"):
        fname = f"{basename_prefix}_{uuid.uuid4().hex[:8]}.tif"
        fpath = os.path.join(self.up_dir, fname)
        height, width = 4, 4
        transform = from_origin(100, 21, 1, 1)
        arr = np.zeros((4, height, width), dtype=np.float32)
        arr[0] = 1
        arr[1] = 1
        arr[2] = 1
        arr[3] = 3

        with rasterio.open(
            fpath,
            "w",
            driver="GTiff",
            height=height,
            width=width,
            count=4,
            dtype=np.float32,
            transform=transform,
            crs="EPSG:4326",
        ) as dst:
            for i in range(4):
                dst.write(arr[i], i + 1)

        self.created_upload_files.append(fpath)
        return fname

    def _post_spectral(self, payload):
        resp = self.client.post("/api/analysis/spectral_indices", json=payload)
        data = json.loads(resp.data)
        return resp.status_code, data

    def _record_result_files(self, result_payload):
        result_urls = result_payload.get("urls", []) if isinstance(result_payload, dict) else result_payload
        for url in result_urls:
            if isinstance(url, str) and url.startswith(generate_url):
                relative = url[len(generate_url):]
                local = os.path.join(self.generate_dir, relative)
                self.created_result_files.append(local)

        latest = Analysis.query.filter_by(type=8).order_by(Analysis.id.desc()).first()
        if latest and latest.before_img and latest.before_img.startswith(generate_url):
            relative = latest.before_img[len(generate_url):]
            local = os.path.join(self.generate_dir, relative)
            self.created_result_files.append(local)

    def test_ndvi_success_and_persisted_metadata(self):
        tif = self._make_tif(4, "ndvi")
        status, body = self._post_spectral(
            {
                "list": [up_url + tif],
                "index_type": "NDVI",
                "year": "2024",
                "band_map": {"nir": 4, "red": 3},
            }
        )

        self.assertEqual(status, 200)
        self.assertEqual(body["code"], 0)
        self.assertEqual(len(body["data"]["urls"]), 1)
        self._record_result_files(body["data"])
        self.assertTrue(os.path.exists(self.created_result_files[0]))

        latest = Analysis.query.filter_by(type=8).order_by(Analysis.id.desc()).first()
        self.assertIsNotNone(latest)
        meta = json.loads(latest.data)
        self.assertEqual(meta["index_type"], "NDVI")
        self.assertEqual(meta["year"], "2024")
        self.assertIn("min", meta)
        self.assertIn("max", meta)
        self.assertIn("mean", meta)

    def test_standalone_index_does_not_load_default_mines_or_sync_global_workbooks(self):
        tif = self._make_constant_geo_tif("standalone")
        with (
            patch(
                "applications.interface.analysis._default_kml_path",
                side_effect=AssertionError("standalone must not load default mines"),
            ),
            patch("applications.interface.analysis.sync_miner_index_rows") as sync_global,
        ):
            status, body = self._post_spectral(
                {
                    "list": [up_url + tif],
                    "index_type": "NDVI",
                    "year": "2024",
                    "band_map": {"nir": 4, "red": 3},
                }
            )

        self.assertEqual(status, 200)
        self.assertEqual(body["code"], 0)
        sync_global.assert_not_called()
        self._record_result_files(body["data"])
        latest = Analysis.query.filter_by(type=8).order_by(Analysis.id.desc()).first()
        meta = json.loads(latest.data)
        self.assertEqual(meta["mean_scope"], "image")
        self.assertEqual(meta["matched_fid_list"], [])

    def test_all_supported_indices(self):
        tif = self._make_tif(7, "allidx")
        for index_type, band_map in (
            ("NDVI", {"nir": 4, "red": 3}),
            ("NDWI", {"green": 2, "nir": 4}),
            ("NDBI", {"swir": 6, "nir": 4}),
            ("NDSI", {"green": 2, "swir": 6}),
        ):
            status, body = self._post_spectral(
                {
                    "list": [up_url + tif],
                    "index_type": index_type,
                    "band_map": band_map,
                }
            )
            self.assertEqual(status, 200)
            self.assertEqual(body["code"], 0, msg=f"{index_type} should succeed")
            self.assertTrue(body["data"]["urls"])
            self._record_result_files(body["data"])

    def test_five_band_defaults_cover_swir_indices(self):
        tif = self._make_tif(5, "fiveband")
        for index_type in ("NDBI", "NDSI"):
            status, body = self._post_spectral(
                {
                    "list": [up_url + tif],
                    "index_type": index_type,
                }
            )
            self.assertEqual(status, 200)
            self.assertEqual(body["code"], 0, msg=f"{index_type} should succeed with 5-band defaults")
            self.assertTrue(body["data"]["urls"])
            self._record_result_files(body["data"])

    def test_invalid_index_type(self):
        tif = self._make_tif(4, "invalid")
        status, body = self._post_spectral(
            {
                "list": [up_url + tif],
                "index_type": "ABC",
            }
        )
        self.assertEqual(status, 200)
        self.assertEqual(body["code"], 1)
        self.assertIn("不支持的指数类型", body["msg"])

    def test_missing_required_band_mapping(self):
        tif = self._make_tif(4, "missingmap")
        status, body = self._post_spectral(
            {
                "list": [up_url + tif],
                "index_type": "NDBI",
                "band_map": {"nir": 4},
            }
        )
        self.assertEqual(status, 200)
        self.assertEqual(body["code"], 1)
        self.assertIn("波段映射缺失", body["msg"])

    def test_band_index_out_of_range(self):
        tif = self._make_tif(4, "range")
        status, body = self._post_spectral(
            {
                "list": [up_url + tif],
                "index_type": "NDVI",
                "band_map": {"nir": 99, "red": 3},
            }
        )
        self.assertEqual(status, 200)
        self.assertEqual(body["code"], 1)
        self.assertIn("波段序号超出范围", body["msg"])

    def test_kml_merge_inserts_and_updates_by_fid(self):
        tmp = tempfile.mkdtemp()
        self.temp_dirs.append(tmp)
        base_kml = os.path.join(tmp, "base.kml")
        incoming_kml = os.path.join(tmp, "incoming.kml")
        with open(base_kml, "w", encoding="utf-8") as f:
            f.write(
                """<?xml version="1.0" encoding="UTF-8"?><kml xmlns="http://www.opengis.net/kml/2.2"><Document>
                <Placemark><name>100</name><Polygon><outerBoundaryIs><LinearRing><coordinates>
                100,20,0 101,20,0 101,21,0 100,21,0 100,20,0
                </coordinates></LinearRing></outerBoundaryIs></Polygon></Placemark>
                </Document></kml>"""
            )
        with open(incoming_kml, "w", encoding="utf-8") as f:
            f.write(
                """<?xml version="1.0" encoding="UTF-8"?><kml xmlns="http://www.opengis.net/kml/2.2"><Document>
                <Placemark><name>100</name><Polygon><outerBoundaryIs><LinearRing><coordinates>
                102,20,0 103,20,0 103,21,0 102,21,0 102,20,0
                </coordinates></LinearRing></outerBoundaryIs></Polygon></Placemark>
                <Placemark><name>101</name><Polygon><outerBoundaryIs><LinearRing><coordinates>
                104,20,0 105,20,0 105,21,0 104,21,0 104,20,0
                </coordinates></LinearRing></outerBoundaryIs></Polygon></Placemark>
                </Document></kml>"""
            )

        summary = merge_kml_increment(Path(base_kml), Path(incoming_kml))
        self.assertEqual(summary["updated"], 1)
        self.assertEqual(summary["inserted"], 1)
        features = load_kml_features(Path(base_kml))
        self.assertEqual([fid for fid, _ in features].count("100"), 1)
        self.assertEqual([fid for fid, _ in features].count("101"), 1)

    def test_sync_miner_index_rows_upserts_year_column(self):
        tmp = tempfile.mkdtemp()
        self.temp_dirs.append(tmp)
        sync_miner_index_rows("NDVI", "2024", [{"fid": 9001, "mean": 0.5}], miner_dir=Path(tmp))
        sync_miner_index_rows("NDVI", "2024", [{"fid": 9001, "mean": 0.6}], miner_dir=Path(tmp))

        wb = load_workbook(os.path.join(tmp, "NDVI_2year.xlsx"), data_only=True)
        ws = wb[wb.sheetnames[0]]
        headers = [cell.value for cell in ws[1]]
        year_col = headers.index("2024") + 1
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][0], 9001)
        self.assertAlmostEqual(ws.cell(row=2, column=year_col).value, 0.6)

    def test_sync_miner_index_rows_can_preserve_existing_cells(self):
        tmp = tempfile.mkdtemp()
        self.temp_dirs.append(tmp)
        sync_miner_index_rows("NDVI", "2024", [{"fid": 9001, "mean": 0.5}], miner_dir=Path(tmp))
        result = sync_miner_index_rows(
            "NDVI",
            "2024",
            [{"fid": 9001, "mean": 0.6}],
            miner_dir=Path(tmp),
            overwrite_existing=False,
        )

        wb = load_workbook(os.path.join(tmp, "NDVI_2year.xlsx"), data_only=True)
        ws = wb[wb.sheetnames[0]]
        headers = [cell.value for cell in ws[1]]
        year_col = headers.index("2024") + 1
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][0], 9001)
        self.assertAlmostEqual(ws.cell(row=2, column=year_col).value, 0.5)
        self.assertEqual(result["skipped_existing"], 1)

    def test_ndvi_mean_uses_matching_kml_polygon_fid(self):
        tif = self._make_constant_geo_tif("ndvi_poly")
        # kml_path 只接受受控 vector 根（miner/）内的 basename 矢量文件
        miner_root = Path(__file__).resolve().parents[1] / "miner"
        kml_name = f"mine_test_{uuid.uuid4().hex[:8]}.kml"
        kml_path = miner_root / kml_name
        with open(kml_path, "w", encoding="utf-8") as f:
            f.write(
                """<?xml version="1.0" encoding="UTF-8"?><kml xmlns="http://www.opengis.net/kml/2.2"><Document>
                <Placemark><name>9002</name><Polygon><outerBoundaryIs><LinearRing><coordinates>
                100,17,0 104,17,0 104,21,0 100,21,0 100,17,0
                </coordinates></LinearRing></outerBoundaryIs></Polygon></Placemark>
                </Document></kml>"""
            )
        self.created_upload_files.append(str(kml_path))

        status, body = self._post_spectral(
            {
                "list": [up_url + tif],
                "index_type": "NDVI",
                "band_map": {"nir": 4, "red": 3},
                "kml_path": kml_name,
            }
        )

        self.assertEqual(status, 200)
        self.assertEqual(body["code"], 0)
        self._record_result_files(body["data"])
        latest = Analysis.query.filter_by(type=8).order_by(Analysis.id.desc()).first()
        meta = json.loads(latest.data)
        self.assertEqual(meta["matched_fid_list"], [9002])
        self.assertAlmostEqual(meta["mean"], 0.5)
        self.assertEqual(meta["fid_stats"][0]["pixel_count"], 16)

    def test_project_index_uses_only_active_project_vector_and_writes_matching_fid(self):
        tif = self._make_constant_geo_tif("project_ndvi")
        project = create_project({"name": "Dali", "region": "Dali"})
        db.session.add(
            ProjectMineBinding(
                project_id=project["id"],
                mine_fid=9002,
                sort_order=0,
            )
        )
        storage_root = Path(tempfile.mkdtemp())
        self.temp_dirs.append(str(storage_root))
        relative = f"projects/{project['id']}/mines/1/mines.geojson"
        vector_path = storage_root / relative
        vector_path.parent.mkdir(parents=True)
        vector_path.write_text(
            json.dumps(
                {
                    "type": "FeatureCollection",
                    "features": [
                        {
                            "type": "Feature",
                            "properties": {"FID_1": 9002},
                            "geometry": {
                                "type": "Polygon",
                                "coordinates": [[[100, 17], [104, 17], [104, 21], [100, 21], [100, 17]]],
                            },
                        }
                    ],
                }
            ),
            encoding="utf-8",
        )
        db.session.add(
            ProjectSpatialResource(
                project_id=project["id"],
                resource_type="mine_vector",
                version=1,
                status="active",
                source_path=relative,
                normalized_path=relative,
                source_format="geojson",
            )
        )
        db.session.commit()
        previous_root = os.environ.get("PROJECT_STORAGE_ROOT")
        os.environ["PROJECT_STORAGE_ROOT"] = str(storage_root)
        try:
            status, body = self._post_spectral(
                {
                    "project_id": project["id"],
                    "list": [{
                        "src": up_url + tif,
                        "raw_tiff_path": os.path.join(self.up_dir, tif),
                    }],
                    "index_type": "NDVI",
                    "year": "2024",
                    "band_map": {"nir": 4, "red": 3},
                }
            )
        finally:
            if previous_root is None:
                os.environ.pop("PROJECT_STORAGE_ROOT", None)
            else:
                os.environ["PROJECT_STORAGE_ROOT"] = previous_root

        self.assertEqual(status, 200)
        self.assertEqual(body["code"], 0)
        self.assertEqual(body["data"]["routing"]["synced_fids"], [9002])
        output = storage_root / f"projects/{project['id']}/outputs/indices/9002.json"
        payload = json.loads(output.read_text(encoding="utf-8"))
        self.assertEqual(payload["ndvi"]["data"], [{"year": 2024, "value": 0.5}])


if __name__ == "__main__":
    unittest.main()
