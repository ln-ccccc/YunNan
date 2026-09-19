#!/usr/bin/env python3
import json
import os
import re
import shutil
from pathlib import Path

from openpyxl import load_workbook

from applications import create_app
from applications.extensions import db
from applications.models.project import Project, ProjectActivityLog, ProjectDataset
from applications.models.project_spatial import ProjectSpatialJob, ProjectSpatialResource
from applications.project_hub.spatial_service import import_mine_vector, list_basemap_candidates, register_basemap
from applications.project_hub.yunnan_seed import PROJECT_REMARK


INDEX_FILES = {
    "ndvi": "NDVI_2year.xlsx",
    "ndbi": "NDBI_by_fid_2year_avg.xlsx",
    "ndwi": "NDWI_by_fid_2year_avg.xlsx",
    "ndsi": "NDSI_by_fid_2year_avg.xlsx",
}


def _index_stats(records):
    if not records:
        return {"mean": 0, "trend": 0, "mk_trend": "no_data"}
    values = [item["value"] for item in records]
    years = [item["year"] for item in records]
    mean = sum(values) / len(values)
    if len(values) < 2:
        return {"mean": round(mean, 3), "trend": 0, "mk_trend": "insufficient_data"}
    count = len(values)
    denominator = count * sum(year * year for year in years) - sum(years) ** 2
    slope = 0 if denominator == 0 else (
        count * sum(year * value for year, value in zip(years, values)) - sum(years) * sum(values)
    ) / denominator
    trend = "upward" if slope > 0.0005 else ("downward" if slope < -0.0005 else "stable")
    return {"mean": round(mean, 3), "trend": round(slope, 5), "mk_trend": trend}


def _read_index_workbook(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = workbook.active.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    fid_index = next((index for index, value in enumerate(headers) if value.casefold() in {"fid", "fid_1"}), None)
    year_columns = []
    for index, header in enumerate(headers):
        match = re.search(r"(?:19|20)\d{2}", header)
        if match:
            year_columns.append((index, int(match.group(0))))
    result = {}
    if fid_index is None:
        return result
    for row in rows:
        try:
            fid = int(row[fid_index])
        except (TypeError, ValueError):
            continue
        records = []
        for column_index, year in year_columns:
            try:
                value = float(row[column_index])
            except (TypeError, ValueError):
                continue
            records.append({"year": year, "value": value})
        if records:
            result[fid] = records
    return result


def migrate_indices(project):
    storage_root = Path(os.environ.get("PROJECT_STORAGE_ROOT", "/project_storage")).resolve()
    output_dir = storage_root / "projects" / str(project.id) / "outputs" / "indices"
    source_dir = storage_root / "projects" / str(project.id) / "outputs" / "index_sources"
    marker = output_dir / ".migrated-v2"
    if marker.exists():
        return "existing"
    # 源 xlsx 全部缺失（未挂载/交付包漏装）时不迁移也不写完成标记：
    # 否则失败被固化为"已完成"，补装文件后也不会重试（2026-09-20 审查 P2）
    available_sources = {
        index_name: source
        for index_name, filename in INDEX_FILES.items()
        if (source := Path("/app/miner") / filename).is_file()
    }
    if not available_sources:
        raise FileNotFoundError(
            "指数迁移源缺失：/app/miner 下未找到任何 "
            f"{sorted(INDEX_FILES.values())}，跳过迁移且不写完成标记"
        )
    output_dir.mkdir(parents=True, exist_ok=True)
    source_dir.mkdir(parents=True, exist_ok=True)
    for old_payload in output_dir.glob("*.json"):
        old_payload.unlink()
    bound_fids = {row.mine_fid for row in project.mines}
    by_fid = {}
    for index_name, filename in INDEX_FILES.items():
        source = available_sources.get(index_name)
        if source is None:
            continue
        target = source_dir / filename
        if not target.exists():
            shutil.copy2(source, target)
        dataset = ProjectDataset.query.filter_by(
            project_id=project.id,
            dataset_kind="mine_indices",
            display_name=filename,
        ).first()
        if dataset is None:
            db.session.add(
                ProjectDataset(
                    project_id=project.id,
                    dataset_kind="mine_indices",
                    display_name=filename,
                    file_path=target.relative_to(storage_root).as_posix(),
                    source_format="xlsx",
                )
            )
        for fid, records in _read_index_workbook(source).items():
            if fid not in bound_fids:
                continue
            by_fid.setdefault(fid, {})[index_name] = records
    for fid, series_by_index in by_fid.items():
        payload = {"fid": fid}
        for index_name, filename in INDEX_FILES.items():
            records = series_by_index.get(index_name, [])
            payload[index_name] = {
                "data": records,
                **_index_stats(records),
                "available": bool(records),
                "source_file": filename,
                "reason": None if records else "missing_mine_data",
                "message": "" if records else "暂无项目数据",
            }
        (output_dir / f"{fid}.json").write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    marker.write_text("ok\n", encoding="utf-8")
    db.session.add(
        ProjectActivityLog(
            project_id=project.id,
            event_type="legacy_indices_migrated",
            payload_json=json.dumps({"mine_count": len(by_fid), "files": list(INDEX_FILES.values())}, ensure_ascii=False),
        )
    )
    db.session.commit()
    return len(by_fid)


def migrate(kml_path):
    project = Project.query.filter_by(remark=PROJECT_REMARK, deleted_at=None).first()
    if project is None:
        raise RuntimeError("未找到云南种子项目")
    result = {"project_id": project.id, "mine_resource": "existing", "basemap_job": "existing"}
    mine_resource = ProjectSpatialResource.query.filter_by(
        project_id=project.id,
        resource_type="mine_vector",
        status="active",
    ).first()
    if mine_resource is None:
        imported = import_mine_vector(
            project.id,
            Path(kml_path).name,
            Path(kml_path).read_text(encoding="utf-8"),
            {"fid": "FID_1"},
        )
        result["mine_resource"] = imported["resource"]["id"]
    result["indices"] = migrate_indices(project)

    basemap = ProjectSpatialResource.query.filter_by(
        project_id=project.id,
        resource_type="basemap",
        status="active",
    ).first()
    pending_job = (
        ProjectSpatialJob.query.filter_by(project_id=project.id, job_type="basemap_tiles")
        .filter(ProjectSpatialJob.status.in_(["queued", "running"]))
        .first()
    )
    if basemap is None and pending_job is None:
        candidates = list_basemap_candidates(project.id)
        preferred = next((item for item in candidates if "Level_15" in item["filename"]), None)
        if preferred is None and len(candidates) == 1:
            preferred = candidates[0]
        if preferred is None:
            result["basemap_job"] = "waiting_for_candidate"
        else:
            queued = register_basemap(project.id, preferred["candidate"], 8, 15)
            result["basemap_job"] = queued["job"]["id"]
    return result


def main():
    app = create_app(os.getenv("FLASK_CONFIG", "development"))
    with app.app_context():
        result = migrate(os.getenv("YUNNAN_KML_PATH", "/app/miner/yunnan.kml"))
        print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
