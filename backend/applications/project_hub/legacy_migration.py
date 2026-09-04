import json
import re
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from sqlalchemy import desc

from applications.extensions import db
from applications.kml_roi.index_sync import INDEX_FILE_MAP, sync_miner_index_rows
from applications.models.analysis import Analysis
from applications.models.project import Project, ProjectActivityLog, ProjectDataset, ProjectMineBinding


KML_NS = {"kml": "http://www.opengis.net/kml/2.2"}
YEAR_RE = re.compile(r"(20\d{2})")
MASK_YEAR_RE = re.compile(r"\+(20\d{2})_mask\.png$")
LEGACY_PROJECT_REMARK = "legacy_migration:auto"
MAX_DB_INT = 2147483647


def _repo_root() -> Path:
    return Path(__file__).resolve().parents[3]


def _default_miner_root() -> Path:
    return _repo_root() / "miner"


def _default_output_root() -> Path:
    return _default_miner_root() / "change_matrix_outputs"


def _default_kml_path() -> Path:
    return _default_miner_root() / "yunnan.kml"


def _default_static_root() -> Path:
    return _repo_root() / "backend" / "static"


def _safe_json_load(value, default=None):
    if value in (None, ""):
        return {} if default is None else default
    if isinstance(value, (dict, list)):
        return value
    try:
        return json.loads(value)
    except Exception:
        return {} if default is None else default


def _json_dump(value):
    return json.dumps(value or {}, ensure_ascii=False)


def _to_int(value):
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return int(float(text))
    except Exception:
        return None


def _to_float(value):
    text = str(value or "").strip()
    if not text:
        return None


def _valid_mine_fid(fid):
    return isinstance(fid, int) and 0 < fid <= MAX_DB_INT
    try:
        return float(text)
    except Exception:
        return None


def _find_or_create_project(project_name: str, manager: str, min_year, max_year):
    project = Project.query.filter_by(name=project_name, deleted_at=None).first()
    project_created = False
    if project is None:
        project = Project(
            name=project_name,
            region="云南省",
            manager=manager or "admin",
            remark=LEGACY_PROJECT_REMARK,
            status="active",
            monitor_start_year=min_year,
            monitor_end_year=max_year,
        )
        db.session.add(project)
        db.session.flush()
        project_created = True
        return project, project_created

    project.deleted_at = None
    if not project.region:
        project.region = "云南省"
    if manager and not project.manager:
        project.manager = manager
    if project.status == "draft":
        project.status = "active"
    if project.remark in (None, ""):
        project.remark = LEGACY_PROJECT_REMARK
    if min_year is not None and (project.monitor_start_year is None or min_year < project.monitor_start_year):
        project.monitor_start_year = min_year
    if max_year is not None and (project.monitor_end_year is None or max_year > project.monitor_end_year):
        project.monitor_end_year = max_year
    return project, project_created


def _load_kml_snapshots(kml_path: Path):
    snapshots = {}
    if not kml_path.exists():
        return snapshots
    root = ET.parse(kml_path).getroot()
    for placemark in root.findall(".//kml:Placemark", KML_NS):
        fields = {}
        for item in placemark.findall(".//kml:SimpleData", KML_NS):
            key = str(item.attrib.get("name") or "").strip()
            fields[key] = "".join(item.itertext()).strip()

        fid = _to_int(fields.get("FID_1") or placemark.findtext("kml:name", default="", namespaces=KML_NS))
        if not _valid_mine_fid(fid):
            continue

        mine_name = (
            fields.get("GGKSMC")
            or fields.get("SBKSMC")
            or fields.get("ZLKSMC")
            or placemark.findtext("kml:name", default="", namespaces=KML_NS).strip()
            or f"矿山 {fid}"
        )
        snapshots[fid] = {
            "mine_name_snapshot": mine_name,
            "city_snapshot": fields.get("SHI") or fields.get("SHI_1") or "",
            "area_snapshot": _to_float(fields.get("TBTYMJ_1") or fields.get("TBTYMJ") or fields.get("SHAPE_Area")),
            "status_snapshot": fields.get("HFZLQK") or fields.get("ZLHFZLQK") or "",
        }
    return snapshots


def _iter_change_output_entries(output_root: Path):
    entries = []
    if not output_root.exists():
        return entries

    def sort_key(path_obj: Path):
        fid = _to_int(path_obj.name)
        return (fid is None, fid if fid is not None else path_obj.name)

    for child in sorted(output_root.iterdir(), key=sort_key):
        if not child.is_dir():
            continue
        fid = _to_int(child.name)
        if not _valid_mine_fid(fid):
            continue
        file_names = sorted(item.name for item in child.iterdir() if item.is_file())
        if not file_names:
            continue
        years = sorted(
            {
                int(match.group(1))
                for name in file_names
                for match in [MASK_YEAR_RE.search(name)]
                if match
            }
        )
        entries.append(
            {
                "fid": fid,
                "path": child,
                "available_years": years,
                "file_names": file_names,
                "has_change_matrix_percent": "change_matrix_percent_rownorm.csv" in file_names,
                "has_class_ratio": "class_ratio_percent.json" in file_names,
            }
        )
    return entries


def _workbook_summary(path_obj: Path):
    if not path_obj.exists():
        return {
            "path": path_obj,
            "mine_fids": [],
            "available_years": [],
        }
    workbook = load_workbook(path_obj, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    years = sorted(
        {
            int(match.group(1))
            for value in header_row
            for match in [YEAR_RE.search(str(value or ""))]
            if match
        }
    )
    mine_fids = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        fid = _to_int(row[0] if row else None)
        if _valid_mine_fid(fid):
            mine_fids.append(fid)
    return {
        "path": path_obj,
        "mine_fids": sorted(set(mine_fids)),
        "available_years": years,
    }


def _resolve_analysis_file_path(raw_path: str, static_root: Path):
    text = str(raw_path or "").strip()
    if not text:
        return text
    if text.startswith("/_uploads/photos/"):
        relative = text[len("/_uploads/photos/"):]
        candidate = static_root / "upload" / relative
        return str(candidate) if candidate.exists() else text
    if text.startswith("static/"):
        candidate = static_root.parent / text
        return str(candidate) if candidate.exists() else text
    return text


def _load_analysis_rows():
    return Analysis.query.order_by(desc(Analysis.create_time), desc(Analysis.id)).all()


def _sync_analysis_workbooks(analysis_rows, miner_root: Path):
    generated_files = []
    synced_rows = 0
    skipped_existing = 0
    warnings = []

    for row in analysis_rows:
        if row.type != 8:
            continue
        meta = _safe_json_load(row.data, {})
        index_type = str(meta.get("index_type") or "").strip().upper()
        year = str(meta.get("year") or "").strip()
        fid_stats = meta.get("fid_stats") or []
        if not index_type or not fid_stats:
            continue
        target_path = miner_root / INDEX_FILE_MAP.get(index_type, "")
        existed_before = target_path.exists()
        normalized_rows = [
            {"fid": item.get("fid"), "mean": item.get("mean")}
            for item in fid_stats
            if item.get("fid") not in (None, "") and item.get("mean") is not None
        ]
        result = sync_miner_index_rows(
            index_type,
            year,
            normalized_rows,
            miner_dir=miner_root,
            overwrite_existing=False,
        )
        if result.get("synced"):
            synced_rows += int(result.get("rows") or 0)
            skipped_existing += int(result.get("skipped_existing") or 0)
            if not existed_before and index_type not in generated_files:
                generated_files.append(index_type)
        else:
            warnings.append(
                {
                    "analysis_id": row.id,
                    "index_type": index_type,
                    "reason": result.get("reason"),
                }
            )

    return {
        "generated_files": generated_files,
        "synced_rows": synced_rows,
        "skipped_existing": skipped_existing,
        "warnings": warnings,
    }


def _build_workbook_datasets(index_summaries):
    datasets = []
    for index_type, summary in sorted(index_summaries.items()):
        path_obj = summary["path"]
        if not path_obj.exists():
            continue
        years = summary["available_years"]
        datasets.append(
            {
                "dataset_kind": "report",
                "display_name": f"历史 {index_type} 指数时序",
                "file_path": str(path_obj),
                "source_format": "xlsx",
                "mine_fid": None,
                "year_start": min(years) if years else None,
                "year_end": max(years) if years else None,
                "slice_config_json": {
                    "legacy_source": "miner_index_workbook",
                    "index_type": index_type,
                    "mine_count": len(summary["mine_fids"]),
                    "available_years": years,
                },
            }
        )
    return datasets


def _build_change_output_datasets(entries):
    datasets = []
    for entry in entries:
        datasets.append(
            {
                "dataset_kind": "inference_result",
                "display_name": f"历史变化矩阵输出 FID {entry['fid']}",
                "file_path": str(entry["path"]),
                "source_format": "change_matrix_dir",
                "mine_fid": entry["fid"],
                "year_start": min(entry["available_years"]) if entry["available_years"] else None,
                "year_end": max(entry["available_years"]) if entry["available_years"] else None,
                "slice_config_json": {
                    "legacy_source": "change_matrix_outputs",
                    "available_years": entry["available_years"],
                    "file_count": len(entry["file_names"]),
                    "has_change_matrix_percent": entry["has_change_matrix_percent"],
                    "has_class_ratio": entry["has_class_ratio"],
                },
            }
        )
    return datasets


def _build_analysis_datasets(analysis_rows, static_root: Path):
    datasets = []
    for row in analysis_rows:
        meta = _safe_json_load(row.data, {})
        matched_fids = [_to_int(item) for item in (meta.get("matched_fid_list") or [])]
        matched_fids = [item for item in matched_fids if _valid_mine_fid(item)]
        mine_fid = matched_fids[0] if len(matched_fids) == 1 else None
        if row.type == 8:
            index_type = str(meta.get("index_type") or "INDEX").strip().upper()
            year = str(meta.get("year") or "").strip()
            display_name = f"历史 {index_type} 分析记录 #{row.id}"
            if year:
                display_name = f"历史 {index_type} 分析记录 {year} #{row.id}"
            year_start = _to_int(year)
            year_end = _to_int(year)
        else:
            index_type = ""
            display_name = f"历史分析记录 #{row.id}"
            year_start = None
            year_end = None
        datasets.append(
            {
                "dataset_kind": "inference_result",
                "display_name": display_name,
                "file_path": _resolve_analysis_file_path(row.after_img or row.before_img, static_root),
                "source_format": "analysis_record",
                "mine_fid": mine_fid,
                "year_start": year_start,
                "year_end": year_end,
                "slice_config_json": {
                    "legacy_source": "analysis_history",
                    "analysis_id": row.id,
                    "analysis_type": row.type,
                    "index_type": index_type,
                    "matched_fid_list": matched_fids[:200],
                    "raw_after_img": row.after_img,
                    "raw_before_img": row.before_img,
                },
            }
        )
    return datasets


def _collect_mine_fids(change_entries, index_summaries, analysis_rows):
    mine_fids = {entry["fid"] for entry in change_entries}
    for summary in index_summaries.values():
        mine_fids.update(summary["mine_fids"])
    for row in analysis_rows:
        meta = _safe_json_load(row.data, {})
        mine_fids.update(
            fid
            for fid in (_to_int(item) for item in (meta.get("matched_fid_list") or []))
            if _valid_mine_fid(fid)
        )
    return sorted(mine_fids)


def _dataset_key(item):
    return (
        str(item.get("dataset_kind") or "").strip(),
        str(item.get("file_path") or "").strip(),
        _to_int(item.get("mine_fid")) or 0,
    )


def _upsert_mine_bindings(project: Project, mine_fids, snapshots):
    existing = {row.mine_fid: row for row in project.mines}
    added = 0
    for sort_order, fid in enumerate(mine_fids, start=1):
        row = existing.get(fid)
        if row is None:
            row = ProjectMineBinding(project_id=project.id, mine_fid=fid)
            db.session.add(row)
            existing[fid] = row
            added += 1
        snapshot = snapshots.get(fid, {})
        row.mine_name_snapshot = snapshot.get("mine_name_snapshot") or row.mine_name_snapshot or f"矿山 {fid}"
        row.city_snapshot = snapshot.get("city_snapshot") or row.city_snapshot
        if snapshot.get("area_snapshot") is not None:
            row.area_snapshot = snapshot["area_snapshot"]
        row.status_snapshot = snapshot.get("status_snapshot") or row.status_snapshot
        row.sort_order = sort_order
    return added, len(existing)


def _upsert_datasets(project: Project, dataset_items):
    existing = {_dataset_key(item.__dict__): item for item in project.datasets}
    added = 0
    for item in dataset_items:
        key = _dataset_key(item)
        row = existing.get(key)
        if row is None:
            row = ProjectDataset(
                project_id=project.id,
                dataset_kind=item["dataset_kind"],
                display_name=item["display_name"],
                file_path=item["file_path"],
                source_format=item["source_format"],
                mine_fid=item.get("mine_fid"),
                year_start=item.get("year_start"),
                year_end=item.get("year_end"),
                slice_config_json=_json_dump(item.get("slice_config_json")),
            )
            db.session.add(row)
            existing[key] = row
            added += 1
            continue
        row.display_name = item["display_name"]
        row.source_format = item["source_format"]
        row.year_start = item.get("year_start")
        row.year_end = item.get("year_end")
        row.slice_config_json = _json_dump(item.get("slice_config_json"))
    return added, len(existing)


def _append_activity(project_id, event_type, payload, actor):
    db.session.add(
        ProjectActivityLog(
            project_id=project_id,
            event_type=event_type,
            actor=actor or "system",
            payload_json=_json_dump(payload),
        )
    )


def migrate_legacy_project_data(
    project_name: str = "历史成果迁移项目",
    manager: str = "admin",
    miner_root=None,
    output_root=None,
    kml_path=None,
    static_root=None,
):
    miner_root = Path(miner_root) if miner_root else _default_miner_root()
    output_root = Path(output_root) if output_root else _default_output_root()
    kml_path = Path(kml_path) if kml_path else _default_kml_path()
    static_root = Path(static_root) if static_root else _default_static_root()

    analysis_rows = _load_analysis_rows()
    index_sync = _sync_analysis_workbooks(analysis_rows, miner_root)

    index_summaries = {
        index_type: _workbook_summary(miner_root / filename)
        for index_type, filename in INDEX_FILE_MAP.items()
    }
    change_entries = _iter_change_output_entries(output_root)
    snapshots = _load_kml_snapshots(kml_path)
    mine_fids = _collect_mine_fids(change_entries, index_summaries, analysis_rows)

    all_years = []
    for entry in change_entries:
        all_years.extend(entry["available_years"])
    for summary in index_summaries.values():
        all_years.extend(summary["available_years"])
    for row in analysis_rows:
        meta = _safe_json_load(row.data, {})
        year = _to_int(meta.get("year"))
        if year is not None:
            all_years.append(year)
    min_year = min(all_years) if all_years else None
    max_year = max(all_years) if all_years else None

    project, project_created = _find_or_create_project(project_name, manager, min_year, max_year)

    mine_bindings_added, mine_count = _upsert_mine_bindings(project, mine_fids, snapshots)
    dataset_items = []
    dataset_items.extend(_build_workbook_datasets(index_summaries))
    dataset_items.extend(_build_change_output_datasets(change_entries))
    dataset_items.extend(_build_analysis_datasets(analysis_rows, static_root))
    datasets_added, dataset_count = _upsert_datasets(project, dataset_items)

    payload = {
        "project_created": project_created,
        "mine_bindings_added": mine_bindings_added,
        "mine_count": mine_count,
        "datasets_added": datasets_added,
        "dataset_count": dataset_count,
        "index_sync": index_sync,
        "change_output_count": len(change_entries),
        "analysis_record_count": len(analysis_rows),
    }
    _append_activity(project.id, "legacy_data_migrated", payload, actor=manager or "system")
    db.session.commit()

    return {
        "project_id": project.id,
        "project_created": project_created,
        "mine_bindings_added": mine_bindings_added,
        "mine_count": mine_count,
        "datasets_added": datasets_added,
        "dataset_count": dataset_count,
        "index_sync": index_sync,
        "year_range": [min_year, max_year],
    }
