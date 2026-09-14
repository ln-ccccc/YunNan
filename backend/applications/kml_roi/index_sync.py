import time
from pathlib import Path
from typing import Dict, IO, Iterable, List, Optional

from openpyxl import Workbook, load_workbook

try:
    import fcntl
except ImportError:  # Windows 宿主直跑（本服务生产环境为 Linux 容器）
    fcntl = None
try:
    import msvcrt
except ImportError:
    msvcrt = None


INDEX_FILE_MAP = {
    "NDVI": "NDVI_2year.xlsx",
    "NDBI": "NDBI_by_fid_2year_avg.xlsx",
    "NDWI": "NDWI_by_fid_2year_avg.xlsx",
    "NDSI": "NDSI_by_fid_2year_avg.xlsx",
}

# 并发同步同一工作簿时的锁等待上限：拿不到锁返回明确 warning，
# 不允许长时间阻塞请求线程，更不允许静默覆盖另一进程刚写入的年份列。
LOCK_TIMEOUT_SECONDS = 5.0


def _lock_path_for(target_path: Path) -> Path:
    return target_path.parent / f".{target_path.name}.lock"


class _WorkbookLock:
    """目标工作簿的跨进程排他锁（锁文件实现）。

    - Linux：fcntl.flock(LOCK_EX | LOCK_NB)，非阻塞尝试 + 短间隔重试；
    - Windows：msvcrt.locking(LK_NBLCK)，语义等价；
    - 两者都不可用（罕见）：放行并按单进程语义执行，不阻塞。
    锁随文件句柄释放，进程崩溃时由操作系统自动回收，锁文件本身保留不删。
    """

    def __init__(self, path: Path, timeout: float = LOCK_TIMEOUT_SECONDS, poll_interval: float = 0.2):
        self.path = path
        self.timeout = timeout
        self.poll_interval = poll_interval
        self._fh: Optional[IO] = None

    def acquire(self) -> bool:
        deadline = time.monotonic() + self.timeout
        while True:
            self.path.parent.mkdir(parents=True, exist_ok=True)
            fh = open(self.path, "a+")
            try:
                if fcntl is not None:
                    fcntl.flock(fh.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
                elif msvcrt is not None:
                    fh.seek(0)
                    msvcrt.locking(fh.fileno(), msvcrt.LK_NBLCK, 1)
                self._fh = fh
                return True
            except OSError:
                fh.close()
                if time.monotonic() >= deadline:
                    return False
                time.sleep(self.poll_interval)

    def release(self) -> None:
        fh, self._fh = self._fh, None
        if fh is None:
            return
        try:
            if fcntl is not None:
                fcntl.flock(fh.fileno(), fcntl.LOCK_UN)
            elif msvcrt is not None:
                fh.seek(0)
                msvcrt.locking(fh.fileno(), msvcrt.LK_UNLCK, 1)
        except OSError:
            pass
        finally:
            fh.close()


def _default_miner_dir() -> Path:
    return Path(__file__).resolve().parents[3] / "miner"


def _fid_cell_value(fid):
    text = str(fid).strip()
    return int(text) if text.isdigit() else text


def _find_header(headers: List[object], candidates: Iterable[str]) -> Optional[int]:
    lowered = {str(item).strip().lower() for item in candidates}
    for idx, value in enumerate(headers):
        if str(value or "").strip().lower() in lowered:
            return idx + 1
    return None


def sync_miner_index_rows(
    index_type: str,
    year: str,
    rows: List[Dict],
    miner_dir: Optional[Path] = None,
    overwrite_existing: bool = True,
    lock_timeout: float = LOCK_TIMEOUT_SECONDS,
) -> Dict:
    index_key = str(index_type or "").strip().upper()
    if index_key not in INDEX_FILE_MAP:
        return {"synced": False, "reason": "unsupported_index", "rows": 0}

    year_text = str(year or "").strip()
    if not year_text.isdigit() or len(year_text) != 4:
        return {"synced": False, "reason": "missing_year", "rows": 0}

    valid_rows = [
        row for row in rows
        if row.get("fid") not in (None, "") and row.get("mean") is not None
    ]
    if not valid_rows:
        return {"synced": False, "reason": "no_fid_stats", "rows": 0}

    target_dir = Path(miner_dir) if miner_dir else _default_miner_dir()
    target_dir.mkdir(parents=True, exist_ok=True)
    target_path = target_dir / INDEX_FILE_MAP[index_key]

    # 读-改-写全程持锁：只锁 save 挡不住"读到旧数据后覆盖他方写入"的竞态。
    lock = _WorkbookLock(_lock_path_for(target_path), timeout=lock_timeout)
    if not lock.acquire():
        return {
            "synced": False,
            "reason": "workbook_locked",
            "index_type": index_key,
            "year": year_text,
            "path": str(target_path),
            "rows": 0,
        }

    try:
        return _sync_miner_index_rows_locked(
            target_path=target_path,
            index_key=index_key,
            year_text=year_text,
            valid_rows=valid_rows,
            overwrite_existing=overwrite_existing,
        )
    finally:
        lock.release()


def _sync_miner_index_rows_locked(
    *,
    target_path: Path,
    index_key: str,
    year_text: str,
    valid_rows: List[Dict],
    overwrite_existing: bool,
) -> Dict:
    if target_path.exists():
        wb = load_workbook(target_path)
        ws = wb[wb.sheetnames[0]]
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = target_path.stem[:31]
        ws.cell(row=1, column=1, value="FID_1")

    headers = [cell.value for cell in ws[1]]
    fid_col = _find_header(headers, ("fid_1", "fid"))
    if fid_col is None:
        fid_col = 1
        ws.cell(row=1, column=fid_col, value="FID_1")

    headers = [cell.value for cell in ws[1]]
    year_col = _find_header(headers, (year_text,))
    if year_col is None:
        year_col = ws.max_column + 1
        ws.cell(row=1, column=year_col, value=year_text)

    row_by_fid = {}
    for row_idx in range(2, ws.max_row + 1):
        value = ws.cell(row=row_idx, column=fid_col).value
        if value in (None, ""):
            continue
        row_by_fid[str(value).strip()] = row_idx

    updated = 0
    inserted = 0
    skipped_existing = 0
    for item in valid_rows:
        fid_text = str(item["fid"]).strip()
        row_idx = row_by_fid.get(fid_text)
        if row_idx is None:
            row_idx = ws.max_row + 1
            ws.cell(row=row_idx, column=fid_col, value=_fid_cell_value(fid_text))
            row_by_fid[fid_text] = row_idx
            inserted += 1
        else:
            updated += 1
        target_cell = ws.cell(row=row_idx, column=year_col)
        if not overwrite_existing and target_cell.value not in (None, ""):
            skipped_existing += 1
            continue
        target_cell.value = float(item["mean"])

    wb.save(target_path)
    return {
        "synced": True,
        "index_type": index_key,
        "year": year_text,
        "path": str(target_path),
        "rows": len(valid_rows),
        "inserted": inserted,
        "updated": updated,
        "skipped_existing": skipped_existing,
    }
