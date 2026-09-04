import json
import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.append(os.path.join(os.path.dirname(__file__), "."))

from applications import create_app
from applications.extensions import db
from applications.models.analysis import Analysis
from applications.models.project import Project, ProjectActivityLog
from applications.project_hub.legacy_migration import migrate_legacy_project_data


class TestLegacyProjectMigration(unittest.TestCase):
    def setUp(self):
        self.app = create_app("testing")
        self.app.config["PROPAGATE_EXCEPTIONS"] = True
        self.ctx = self.app.app_context()
        self.ctx.push()
        db.create_all()

        self.temp_dir = tempfile.mkdtemp(prefix="legacy-migrate-")
        self.miner_root = Path(self.temp_dir) / "miner"
        self.output_root = self.miner_root / "change_matrix_outputs"
        self.static_root = Path(self.temp_dir) / "backend" / "static"
        self.kml_path = self.miner_root / "yunnan.kml"

        self.output_root.mkdir(parents=True, exist_ok=True)
        (self.static_root / "upload" / "res").mkdir(parents=True, exist_ok=True)

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.ctx.pop()
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def _write_kml(self):
        self.kml_path.write_text(
            """<?xml version="1.0" encoding="UTF-8"?>
<kml xmlns="http://www.opengis.net/kml/2.2">
  <Document>
    <Placemark>
      <name>Mine 101</name>
      <ExtendedData>
        <SchemaData schemaUrl="#yunnan">
          <SimpleData name="FID_1">101</SimpleData>
          <SimpleData name="SHI">City A</SimpleData>
          <SimpleData name="TBTYMJ">123.5</SimpleData>
          <SimpleData name="HFZLQK">treated</SimpleData>
        </SchemaData>
      </ExtendedData>
    </Placemark>
    <Placemark>
      <name>Mine 102</name>
      <ExtendedData>
        <SchemaData schemaUrl="#yunnan">
          <SimpleData name="FID_1">102</SimpleData>
          <SimpleData name="SHI">City B</SimpleData>
          <SimpleData name="TBTYMJ">88.0</SimpleData>
          <SimpleData name="HFZLQK">untreated</SimpleData>
        </SchemaData>
      </ExtendedData>
    </Placemark>
    <Placemark>
      <name>Mine 103</name>
      <ExtendedData>
        <SchemaData schemaUrl="#yunnan">
          <SimpleData name="FID_1">103</SimpleData>
          <SimpleData name="SHI">City C</SimpleData>
          <SimpleData name="TBTYMJ">66.0</SimpleData>
          <SimpleData name="HFZLQK">unknown</SimpleData>
        </SchemaData>
      </ExtendedData>
    </Placemark>
  </Document>
</kml>""",
            encoding="utf-8",
        )

    def _write_workbook(self, path_obj, headers, rows):
        wb = Workbook()
        ws = wb.active
        ws.append(headers)
        for row in rows:
            ws.append(row)
        wb.save(path_obj)

    def _seed_analysis(self):
        rows = [
            Analysis(
                type=8,
                before_img="/_uploads/photos/res/preview_ndbi.png",
                after_img="/_uploads/photos/res/spectral_ndbi_demo.png",
                data=json.dumps(
                    {
                        "index_type": "NDBI",
                        "year": "2022",
                        "matched_fid_list": [102],
                        "fid_stats": [{"fid": 102, "mean": 0.3, "pixel_count": 12}],
                    }
                ),
            ),
            Analysis(
                type=8,
                before_img="/_uploads/photos/res/preview_ndvi.png",
                after_img="/_uploads/photos/res/spectral_ndvi_demo.png",
                data=json.dumps(
                    {
                        "index_type": "NDVI",
                        "year": "2024",
                        "matched_fid_list": [101],
                        "fid_stats": [{"fid": 101, "mean": 0.1, "pixel_count": 20}],
                    }
                ),
            ),
            Analysis(
                type=3,
                before_img="static/upload/source_demo.tif",
                after_img="/_uploads/photos/res/pred_demo.png",
                data="",
            ),
        ]
        db.session.add_all(rows)
        db.session.commit()

    def _seed_change_outputs(self):
        dir_101 = self.output_root / "101"
        dir_101.mkdir(parents=True, exist_ok=True)
        (dir_101 / "101+2020_mask.png").write_bytes(b"mask2020")
        (dir_101 / "101+2024_mask.png").write_bytes(b"mask2024")
        (dir_101 / "change_matrix_percent_rownorm.csv").write_text(
            ",grass,forest\ngrass,0.7,0.3\nforest,0.1,0.9\n",
            encoding="utf-8",
        )
        (dir_101 / "class_ratio_percent.json").write_text(
            json.dumps(
                {
                    "fid": 101,
                    "years": [2020, 2024],
                    "series_percent": {"bareground": [40, 35]},
                }
            ),
            encoding="utf-8",
        )

        dir_102 = self.output_root / "102"
        dir_102.mkdir(parents=True, exist_ok=True)
        (dir_102 / "102_new.png").write_bytes(b"png")

    def test_migrate_legacy_project_data_indexes_workbooks_outputs_and_history(self):
        self._write_kml()
        self._seed_change_outputs()
        self._seed_analysis()
        self._write_workbook(
            self.miner_root / "NDVI_2year.xlsx",
            ["FID_1", "2024"],
            [[101, 0.9], [103, 0.2]],
        )

        result = migrate_legacy_project_data(
            project_name="Legacy Import",
            manager="admin",
            miner_root=self.miner_root,
            output_root=self.output_root,
            kml_path=self.kml_path,
            static_root=self.static_root,
        )

        self.assertTrue(result["project_created"])
        self.assertGreaterEqual(result["dataset_count"], 7)
        self.assertEqual(result["mine_count"], 3)
        self.assertIn("NDBI", result["index_sync"]["generated_files"])

        project = Project.query.filter_by(name="Legacy Import").first()
        self.assertIsNotNone(project)
        self.assertEqual(len(project.mines), 3)
        self.assertTrue(any(item.display_name == "历史 NDVI 指数时序" for item in project.datasets))
        self.assertTrue(any(item.display_name == "历史 NDBI 指数时序" for item in project.datasets))
        self.assertTrue(any(item.source_format == "change_matrix_dir" for item in project.datasets))
        self.assertTrue(any(item.source_format == "analysis_record" for item in project.datasets))

        wb = load_workbook(self.miner_root / "NDVI_2year.xlsx", data_only=True)
        ws = wb[wb.sheetnames[0]]
        headers = [cell.value for cell in ws[1]]
        year_col = headers.index("2024") + 1
        self.assertAlmostEqual(ws.cell(row=2, column=year_col).value, 0.9)

        ndbi_wb = load_workbook(self.miner_root / "NDBI_by_fid_2year_avg.xlsx", data_only=True)
        ndbi_ws = ndbi_wb[ndbi_wb.sheetnames[0]]
        ndbi_headers = [cell.value for cell in ndbi_ws[1]]
        ndbi_year_col = ndbi_headers.index("2022") + 1
        ndbi_rows = list(ndbi_ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(ndbi_rows[0][0], 102)
        self.assertAlmostEqual(ndbi_ws.cell(row=2, column=ndbi_year_col).value, 0.3)

        events = [row.event_type for row in ProjectActivityLog.query.filter_by(project_id=project.id).all()]
        self.assertIn("legacy_data_migrated", events)

    def test_migrate_legacy_project_data_is_idempotent_for_existing_project(self):
        self._write_kml()
        self._seed_change_outputs()
        self._seed_analysis()
        self._write_workbook(
            self.miner_root / "NDVI_2year.xlsx",
            ["FID_1", "2024"],
            [[101, 0.9]],
        )

        first = migrate_legacy_project_data(
            project_name="Legacy Import",
            manager="admin",
            miner_root=self.miner_root,
            output_root=self.output_root,
            kml_path=self.kml_path,
            static_root=self.static_root,
        )
        second = migrate_legacy_project_data(
            project_name="Legacy Import",
            manager="admin",
            miner_root=self.miner_root,
            output_root=self.output_root,
            kml_path=self.kml_path,
            static_root=self.static_root,
        )

        project = Project.query.filter_by(name="Legacy Import").first()
        self.assertIsNotNone(project)
        self.assertFalse(second["project_created"])
        self.assertEqual(first["dataset_count"], second["dataset_count"])
        self.assertEqual(first["mine_count"], second["mine_count"])
        self.assertEqual(len(project.mines), first["mine_count"])
        self.assertEqual(len(project.datasets), first["dataset_count"])


if __name__ == "__main__":
    unittest.main()
